from pathlib import Path

import openpyxl
from docx import Document

from functions.document_loader import load_document_file, load_documents


def test_load_docx(tmp_path: Path):
    path = tmp_path / "sample.docx"
    doc = Document()
    doc.add_paragraph("INSIGHT AI project overview")
    doc.save(path)

    docs = load_document_file(path)

    assert docs == [{
        "filename": "sample.docx",
        "text": "Document filename: sample.docx\n\nINSIGHT AI project overview",
        "type": "docx",
    }]


def test_load_xlsx(tmp_path: Path):
    path = tmp_path / "sample.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Budget"
    ws.append(["Item", "Amount"])
    ws.append(["Hosting", 25])
    wb.save(path)

    docs = load_document_file(path)

    assert len(docs) == 1
    assert docs[0]["filename"] == "sample.xlsx-Budget"
    assert docs[0]["type"] == "xlsx"
    assert "Item: Hosting | Amount: 25" in docs[0]["text"]


def test_unsupported_file_is_reported(tmp_path: Path):
    (tmp_path / "notes.txt").write_text("ignore me", encoding="utf-8")

    docs, warnings = load_documents(tmp_path)

    assert docs == []
    assert warnings == ["Skipped unsupported file: notes.txt"]
